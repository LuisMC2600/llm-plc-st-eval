# -*- coding: utf-8 -*-
"""Construye las fichas de evaluacion manual a ciegas a partir de paquete_fichas.json.
Salida: Ficha_Evaluador1.xlsx (96 respuestas), Ficha_Evaluador2.xlsx (submuestra de 24) y
Respuestas_Evaluador1.html / Respuestas_Evaluador2.html (texto de las respuestas por codigo).
Uso: python construir_fichas.py <paquete_fichas.json> <carpeta_salida>"""
import json, sys, html, os
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.comments import Comment
from openpyxl.utils import get_column_letter as L

P = json.load(open(sys.argv[1], encoding="utf-8")); OUT = sys.argv[2]; os.makedirs(OUT, exist_ok=True)

CLAVE = {
 "C1": {"tipo": "diag", "items": [
   ("E1", "Condición invertida", "Condición invertida: IF NOT habilitado ejecuta la secuencia cuando debería ser IF habilitado."),
   ("E2", "TON sin reinicio", "El TON t1 se llama con IN:=TRUE permanente y nunca se reinicia entre estados, por lo que las transiciones tras la primera son inmediatas."),
   ("E3", "Estado 3 inalcanzable", "El estado 3 (amarillo intermitente) es inalcanzable: ninguna transición asigna estado:=3."),
   ("E4", "parp no conmuta", "parp nunca se inicializa ni se conmuta, así que el intermitente jamás parpadearía.")],
   "extra": "Hallazgo adicional aceptado por la clave: luzV/luzR no se apagan en el estado 3. Cuenta como hallazgo adicional válido, no como error inventado."},
 "C5": {"tipo": "diag", "items": [
   ("E1", "Interlock kY/kD", "Violación de interlock: en el estado 1 se activa kD y kY solo se apaga al entrar al estado 2, por lo que ambos contactores quedan activos simultáneamente durante al menos un ciclo de scan."),
   ("E2", "Tiempo muerto ausente", "Tiempo muerto ausente: el temporizador t2 se declara pero nunca se usa; la conmutación estrella-triángulo ocurre sin los 100 ms exigidos."),
   ("E3", "t1 sin reinicio en paro", "El temporizador t1 se llama con IN := TRUE permanente y no se reinicia en el paro, de modo que tras un paro y un rearranque la transición a triángulo es inmediata."),
   ("E4", "Paro no apaga kM", "btnParo no desactiva kM, así que el motor queda alimentado por el contactor principal después del paro.")],
   "extra": ""},
 "C6": {"tipo": "diag", "items": [
   ("E1", "División entera", "División entera: rawTemp / 27648 se evalúa en aritmética INT y vale 0 para todo rawTemp entre 0 y 27647 (solo en 27648 da 1, es decir 100.0), de modo que tempC queda prácticamente siempre en 0.0; la conversión a REAL se aplica después del daño."),
   ("E2", "Histéresis inexistente", "La histéresis no existe: la lógica conmuta en un único umbral de 75.0 en lugar de activar en 80.0 y desactivar en 75.0, y la rama ELSIF es inalcanzable para valores por encima de 75.0."),
   ("E3", "Falta rawTemp < 0", "La detección de fuera de rango omite la condición rawTemp < 0, así que un valor negativo no genera alarma."),
   ("E4", "Alarma no se repone ni apaga enfriador", "almSensor nunca vuelve a FALSE al regresar la señal al rango y la alarma no fuerza enfriador := FALSE como exige la especificación.")],
   "extra": ""},
 "C2": {"tipo": "req",
   "es": [("Nivel del tanque", "Entrada que indica el nivel del tanque superior (vacío / lleno)."),
          ("Agua en el pozo", "Entrada que indica si el pozo tiene agua (protección de marcha en seco)."),
          ("Botón de paro", "Entrada del botón para apagar la bomba."),
          ("Bomba", "Salida de mando de la bomba."),
          ("Luz de aviso", "Salida de la luz que avisa cuando algo anda mal.")],
   "amb": [("A1", "Digital o analógica", "Niveles como señales digitales (flotadores) o analógicas."),
           ("A2", "Qué es 'tanque vacío'", "Qué significa exactamente 'tanque vacío' (histéresis entre niveles de arranque y paro)."),
           ("A3", "Marcha en seco y rearme", "Tiempo de protección por marcha en seco y si el rearme es manual o automático."),
           ("A4", "Tipo de botón", "El 'botón para apagarla': ¿paro momentáneo o inhibición retenida?"),
           ("A5", "Qué agrupa la luz", "Qué condiciones agrupa la luz de falla."),
           ("A6", "Retorno del agua", "Comportamiento al recuperarse el agua del pozo.")]},
}
ARIAL = "Arial"
F = lambda **k: Font(name=ARIAL, **({"size": 10} | k))
AMAR = PatternFill("solid", fgColor="FFF2CC"); GRIS = PatternFill("solid", fgColor="E7E6E6"); AZUL = PatternFill("solid", fgColor="1F4E78")
CALC = PatternFill("solid", fgColor="DDEBF7"); fino = Side(style="thin", color="BFBFBF"); BORDE = Border(left=fino, right=fino, top=fino, bottom=fino)
CEN = Alignment(horizontal="center", vertical="center", wrap_text=True); IZQ = Alignment(horizontal="left", vertical="top", wrap_text=True)
FILA_ENC, FILA_EJ, FILA_1 = 4, 5, 6

def cab(ws, col, texto, ancho, nota=None, calc=False):
    c = ws.cell(row=FILA_ENC, column=col, value=texto); c.font = F(bold=True, color="FFFFFF"); c.fill = AZUL; c.alignment = CEN; c.border = BORDE
    ws.column_dimensions[L(col)].width = ancho
    if nota:
        cm = Comment(nota, "Ficha"); cm.width = 380; cm.height = 140; c.comment = cm

def hoja_tarea(wb, t, resp):
    k = CLAVE[t]; ws = wb.create_sheet(t); n = len(resp); ult = FILA_1 + n - 1
    ws["A1"] = "%s — %s" % (t, P["tareas"][t]["titulo"]); ws["A1"].font = F(bold=True, size=13)
    ws["A2"] = ("Rellena solo las celdas amarillas. S = sí, N = no. Lee cada respuesta por su código en el archivo HTML. "
                "La fila EJEMPLO muestra el formato y no se analiza."); ws["A2"].font = F(italic=True)
    ws.row_dimensions[FILA_ENC].height = 62
    cab(ws, 1, "Código", 11)
    sn, num, lista = [], [], []
    if k["tipo"] == "diag":
        col = 2
        for ide, corto, largo in k["items"]:
            cab(ws, col, "%s detectado\n%s" % (ide, corto), 15, "%s — %s" % (ide, largo)); sn.append(col)
            cab(ws, col + 1, "%s corrección válida" % ide, 13, "Marca S si la corrección propuesta para %s resolvería el defecto. Déjala en blanco si no lo detectó." % ide); sn.append(col + 1)
            col += 2
        c_ad, c_inv, c_obs = col, col + 1, col + 2
        cab(ws, c_ad, "Hallazgos adicionales válidos (n.º)", 14, "Defectos reales del código que la respuesta señala y que no están en la clave. " + k["extra"]); num.append(c_ad)
        cab(ws, c_inv, "Errores inventados (n.º)", 13, "Cosas que la respuesta señala como error y no lo son (el código es correcto en ese punto, o contradice la especificación)."); num.append(c_inv)
        cab(ws, c_obs, "Observaciones", 38)
        c1, c2, c3 = c_obs + 1, c_obs + 2, c_obs + 3
        cab(ws, c1, "Detectados (0–4)", 12); cab(ws, c2, "Correcciones válidas (0–4)", 13); cab(ws, c3, "Estado", 11)
        det = [2, 4, 6, 8]; cor = [3, 5, 7, 9]; oblig = det + [c_ad, c_inv]
        ejemplo = {2: "S", 3: "S", 4: "S", 5: "N", 6: "N", 8: "S", 9: "S", c_ad: 1, c_inv: 0, c_obs: "E2: propone reiniciar el TON pero con IN siempre TRUE."}
        for r in range(FILA_EJ, ult + 1):
            ws.cell(row=r, column=c1, value="=" + "+".join('COUNTIF(%s%d,"S")' % (L(c), r) for c in det))
            ws.cell(row=r, column=c2, value="=" + "+".join('COUNTIF(%s%d,"S")' % (L(c), r) for c in cor))
            ws.cell(row=r, column=c3, value='=IF(AND(%s),"Completa","Falta")' % ",".join('%s%d<>""' % (L(c), r) for c in oblig))
        calc = [c1, c2, c3]; ultima_col = c3; estado = c3
    else:
        col = 2
        for corto, largo in k["es"]:
            cab(ws, col, "E/S presente\n%s" % corto, 13, largo + " Marca S si la tabla de E/S la incluye, con cualquier nombre."); sn.append(col); col += 1
        c_tip = col; cab(ws, c_tip, "Tipos IEC válidos en toda la tabla", 13, "S si todos los tipos de dato de la tabla son tipos IEC 61131-3 válidos y coherentes con la señal."); sn.append(c_tip)
        c_pou = col + 1; cab(ws, c_pou, "POU recomendado", 17, "El tipo de POU que la respuesta recomienda."); lista.append(c_pou)
        c_jus = col + 2; cab(ws, c_jus, "Justificación coherente", 13, "S si la justificación del tipo de POU es técnicamente coherente."); sn.append(c_jus)
        col = c_jus + 1; amb = []
        for ide, corto, largo in k["amb"]:
            cab(ws, col, "%s\n%s" % (ide, corto), 13, "%s — %s Marca S si la respuesta plantea esta ambigüedad, con cualquier redacción." % (ide, largo)); sn.append(col); amb.append(col); col += 1
        c_ad = col; cab(ws, c_ad, "Ambigüedades adicionales válidas (n.º)", 14, "Otras preguntas pertinentes al cliente que no están en la clave."); num.append(c_ad)
        c_cod = col + 1; cab(ws, c_cod, "¿Incluyó código?", 11, "El enunciado pide no escribir código. Marca S si la respuesta incluye código ST."); sn.append(c_cod)
        c_obs = col + 2; cab(ws, c_obs, "Observaciones", 38)
        c1, c2, c3 = c_obs + 1, c_obs + 2, c_obs + 3
        cab(ws, c1, "E/S presentes (0–5)", 12); cab(ws, c2, "Ambigüedades de la clave (0–6)", 14); cab(ws, c3, "Estado", 11)
        es = [2, 3, 4, 5, 6]; oblig = es + [c_tip, c_pou, c_jus] + amb + [c_ad, c_cod]
        ejemplo = {2: "S", 3: "S", 4: "S", 5: "S", 6: "S", c_tip: "S", c_pou: "FUNCTION_BLOCK", c_jus: "S", amb[0]: "S", amb[1]: "N", amb[2]: "S", amb[3]: "S", amb[4]: "N", amb[5]: "N", c_ad: 2, c_cod: "N", c_obs: "Pregunta además por la tensión de la bomba."}
        for r in range(FILA_EJ, ult + 1):
            ws.cell(row=r, column=c1, value="=" + "+".join('COUNTIF(%s%d,"S")' % (L(c), r) for c in es))
            ws.cell(row=r, column=c2, value="=" + "+".join('COUNTIF(%s%d,"S")' % (L(c), r) for c in amb))
            ws.cell(row=r, column=c3, value='=IF(AND(%s),"Completa","Falta")' % ",".join('%s%d<>""' % (L(c), r) for c in oblig))
        calc = [c1, c2, c3]; ultima_col = c3; estado = c3
    # filas
    ws.cell(row=FILA_EJ, column=1, value="EJEMPLO")
    for i, x in enumerate(resp): ws.cell(row=FILA_1 + i, column=1, value=x["codigo"])
    entradas = set(sn + num + lista + [c_obs])
    for r in range(FILA_EJ, ult + 1):
        ws.row_dimensions[r].height = 22
        for c in range(1, ultima_col + 1):
            cel = ws.cell(row=r, column=c); cel.border = BORDE; cel.font = F(italic=(r == FILA_EJ), bold=(c == 1 and r != FILA_EJ))
            cel.alignment = IZQ if c == c_obs else Alignment(horizontal="center", vertical="center")
            if r == FILA_EJ: cel.fill = GRIS
            elif c in entradas: cel.fill = AMAR
            elif c in calc: cel.fill = CALC
        if r == FILA_EJ:
            for c, v in ejemplo.items(): ws.cell(row=r, column=c, value=v)
    def rango(c): return "%s%d:%s%d" % (L(c), FILA_EJ, L(c), ult)
    dv = DataValidation(type="list", formula1='"S,N"', allow_blank=True, showErrorMessage=True, errorTitle="Valor no válido", error="Escribe S o N.")
    dn = DataValidation(type="whole", operator="greaterThanOrEqual", formula1="0", allow_blank=True, showErrorMessage=True, errorTitle="Valor no válido", error="Escribe un número entero, 0 o mayor.")
    ws.add_data_validation(dv); ws.add_data_validation(dn)
    for c in sn: dv.add(rango(c))
    for c in num: dn.add(rango(c))
    if lista:
        dl = DataValidation(type="list", formula1='"FUNCTION,FUNCTION_BLOCK,PROGRAM,Varios o ninguno"', allow_blank=True, showErrorMessage=True, errorTitle="Valor no válido", error="Elige una opción de la lista.")
        ws.add_data_validation(dl)
        for c in lista: dl.add(rango(c))
    ws.freeze_panes = ws.cell(row=FILA_1, column=2)
    return (L(estado), FILA_1, ult, n)

def libro(nombre_eval, solo_sub, archivo, html_nombre):
    wb = Workbook(); ins = wb.active; ins.title = "Instrucciones"
    info = {}
    for t in ["C1", "C5", "C6", "C2"]:
        resp = [x for x in P["tareas"][t]["respuestas"] if (x["evaluador2"] or not solo_sub)]
        info[t] = (hoja_tarea(wb, t, resp), resp)
    total = sum(len(v[1]) for v in info.values())
    L1 = [
     ("Evaluación manual a ciegas — %s" % nombre_eval, "t"),
     ("Tesis: LLM como asistentes de programación PLC en la academia (UTP). Tareas complementarias C1, C5, C6 (diagnóstico de código) y C2 (de requisito ambiguo a especificación).", ""),
     ("", ""),
     ("Qué hay que hacer", "h"),
     ("1. Abre el archivo %s. Contiene el enunciado de cada tarea, su clave de evaluación y las %d respuestas, identificadas solo por un código." % (html_nombre, total), ""),
     ("2. En este libro hay una hoja por tarea (C1, C5, C6, C2). Rellena ÚNICAMENTE las celdas amarillas. Las celdas azules se calculan solas.", ""),
     ("3. Califica una tarea completa antes de pasar a la siguiente, en el orden en que aparecen los códigos.", ""),
     ("4. Cada pregunta se responde con S (sí) o N (no). Las columnas de conteo llevan un número entero (0 si no hay ninguno).", ""),
     ("5. La fila gris EJEMPLO solo muestra el formato. No se analiza y no hace falta borrarla.", ""),
     ("6. Al terminar, la tabla de avance de abajo debe mostrar todas las tareas completas. Guarda el archivo con el mismo nombre.", ""),
     ("", ""),
     ("Reglas de calificación", "h"),
     ("Detectado (S): la respuesta identifica el mismo defecto de la clave, con cualquier redacción y aunque lo ubique en otra línea. No cuenta una mención vaga que no señale el defecto concreto.", ""),
     ("Ítems de la clave con dos partes (por ejemplo C6-E4): se marca S si la respuesta identifica al menos una de las partes; anota en Observaciones cuál faltó.", ""),
     ("Corrección válida (S): la corrección propuesta resolvería el defecto. Si el defecto no fue detectado, deja la celda en blanco.", ""),
     ("Hallazgo adicional válido: un defecto real del código que no está en la clave. Error inventado: algo que la respuesta presenta como error y no lo es. Si un mismo punto se repite, cuéntalo una sola vez.", ""),
     ("Califica lo que la respuesta dice, no lo que probablemente quiso decir. Si una respuesta está cortada, califica lo que alcanzó a entregar.", ""),
     ("Las pautas de cada columna están en el comentario del encabezado (pasa el cursor sobre la celda azul oscuro).", ""),
     ("", ""),
     ("Evaluación a ciegas", "h"),
     ("Los códigos no indican qué modelo ni qué estrategia produjo cada respuesta, y el orden está barajado. No consultes el archivo _SELLADO_asignacion_codigos.csv ni la carpeta de datos hasta haber terminado y entregado la ficha.", ""),
     ("Los evaluadores califican por separado y no comentan casos concretos entre sí hasta que ambos hayan terminado.", ""),
     ("", ""),
     ("Avance", "h")]
    for i, (txt, est) in enumerate(L1, 1):
        c = ins.cell(row=i, column=1, value=txt); c.alignment = Alignment(wrap_text=True, vertical="top")
        c.font = F(bold=True, size=14) if est == "t" else (F(bold=True, size=11, color="1F4E78") if est == "h" else F())
        ins.merge_cells(start_row=i, start_column=1, end_row=i, end_column=4)
        if len(txt) > 110: ins.row_dimensions[i].height = 30
    f0 = len(L1) + 1
    for j, h in enumerate(["Tarea", "Respuestas", "Completas", "Pendientes"], 1):
        c = ins.cell(row=f0, column=j, value=h); c.font = F(bold=True, color="FFFFFF"); c.fill = AZUL; c.alignment = CEN; c.border = BORDE
    for i, t in enumerate(["C1", "C5", "C6", "C2"], 1):
        (col, a, b, n), _ = info[t]; r = f0 + i
        ins.cell(row=r, column=1, value=t); ins.cell(row=r, column=2, value=n)
        ins.cell(row=r, column=3, value='=COUNTIF(%s!%s%d:%s%d,"Completa")' % (t, col, a, col, b)); ins.cell(row=r, column=4, value="=B%d-C%d" % (r, r))
        for j in range(1, 5):
            c = ins.cell(row=r, column=j); c.border = BORDE; c.font = F(); c.alignment = CEN
            if j >= 3: c.fill = CALC
    r = f0 + 5
    ins.cell(row=r, column=1, value="Total"); ins.cell(row=r, column=2, value="=SUM(B%d:B%d)" % (f0 + 1, f0 + 4)); ins.cell(row=r, column=3, value="=SUM(C%d:C%d)" % (f0 + 1, f0 + 4)); ins.cell(row=r, column=4, value="=SUM(D%d:D%d)" % (f0 + 1, f0 + 4))
    for j in range(1, 5):
        c = ins.cell(row=r, column=j); c.border = BORDE; c.font = F(bold=True); c.alignment = CEN; c.fill = CALC
    ins.column_dimensions["A"].width = 46; ins.column_dimensions["B"].width = 22; ins.column_dimensions["C"].width = 22; ins.column_dimensions["D"].width = 22
    ins.cell(row=r + 2, column=1, value="Nombre del evaluador:").font = F(bold=True); ins.cell(row=r + 2, column=2).fill = AMAR
    ins.cell(row=r + 3, column=1, value="Fecha de calificación:").font = F(bold=True); ins.cell(row=r + 3, column=2).fill = AMAR
    # Clave
    cl = wb.create_sheet("Clave", 1)
    cl["A1"] = "Clave de evaluación (la misma que acompaña a cada tarea en el archivo HTML)"; cl["A1"].font = F(bold=True, size=13)
    cl["A2"] = "Fuente: clave de evaluación definida con el banco de tareas (src/generar_prompts.py), que nunca se envió a los modelos. Los ítems de E/S y de POU de C2 se derivan del propio requisito del cliente."; cl["A2"].font = F(italic=True); cl["A2"].alignment = Alignment(wrap_text=True); cl.merge_cells("A2:C2"); cl.row_dimensions[2].height = 30
    r = 4
    for j, h in enumerate(["Tarea", "Ítem", "Descripción"], 1):
        c = cl.cell(row=r, column=j, value=h); c.font = F(bold=True, color="FFFFFF"); c.fill = AZUL; c.alignment = CEN; c.border = BORDE
    r += 1
    for t in ["C1", "C5", "C6", "C2"]:
        k = CLAVE[t]
        filas = [(i, l) for i, _, l in k["items"]] + ([("Extra", k["extra"])] if k.get("extra") else []) if k["tipo"] == "diag" else \
                [("E/S: " + c, l) for c, l in k["es"]] + [(i, l) for i, _, l in k["amb"]]
        for ide, largo in filas:
            for j, v in enumerate([t, ide, largo], 1):
                c = cl.cell(row=r, column=j, value=v); c.font = F(bold=(j < 3)); c.border = BORDE; c.alignment = IZQ if j == 3 else Alignment(horizontal="center", vertical="top", wrap_text=True)
            cl.row_dimensions[r].height = 15 * max(1, -(-len(largo) // 95)) + 4; r += 1
    cl.column_dimensions["A"].width = 8; cl.column_dimensions["B"].width = 22; cl.column_dimensions["C"].width = 110
    wb.save(os.path.join(OUT, archivo))
    # HTML
    css = """body{font-family:Arial,Helvetica,sans-serif;margin:0;background:#f4f5f7;color:#1c1e21;line-height:1.45}
header{background:#1f4e78;color:#fff;padding:18px 28px}header h1{margin:0;font-size:20px}header p{margin:6px 0 0;font-size:13px;opacity:.9}
nav{position:sticky;top:0;background:#fff;border-bottom:1px solid #d0d4da;padding:10px 28px;font-size:14px;z-index:5}nav a{margin-right:18px;color:#1f4e78;font-weight:bold;text-decoration:none}
main{max-width:1000px;margin:0 auto;padding:10px 28px 60px}h2{color:#1f4e78;border-bottom:2px solid #1f4e78;padding-bottom:4px;margin-top:38px}
h3{margin:0;font-size:16px}.caja{background:#fff;border:1px solid #d0d4da;border-radius:6px;padding:14px 18px;margin:14px 0}
.clave{background:#fff8e1;border-color:#e6c65c}.resp{border-left:5px solid #1f4e78}.cod{display:flex;justify-content:space-between;align-items:baseline;margin-bottom:8px}
.cod span{font-size:12px;color:#5f6670}pre{white-space:pre-wrap;word-wrap:break-word;font-family:Consolas,'Courier New',monospace;font-size:13px;margin:0}
details summary{cursor:pointer;font-weight:bold;color:#1f4e78}li{margin-bottom:6px}.idx a{display:inline-block;margin:2px 8px 2px 0;font-size:13px;color:#1f4e78}
@media print{nav{position:static}.caja{break-inside:avoid}}"""
    h = ['<!doctype html><html lang="es"><head><meta charset="utf-8"><meta name="viewport" content="width=device-width,initial-scale=1"><title>Respuestas a ciegas — %s</title><style>%s</style></head><body>' % (html.escape(nombre_eval), css)]
    h.append('<header><h1>Respuestas a ciegas — %s</h1><p>%d respuestas identificadas solo por código. Califícalas en %s, tarea por tarea y en este orden.</p></header>' % (html.escape(nombre_eval), total, html.escape(archivo)))
    h.append("<nav>" + "".join('<a href="#%s">%s (%d)</a>' % (t, t, len(info[t][1])) for t in ["C1", "C5", "C6", "C2"]) + "</nav><main>")
    for t in ["C1", "C5", "C6", "C2"]:
        k = CLAVE[t]; resp = info[t][1]
        h.append('<h2 id="%s">%s — %s</h2>' % (t, t, html.escape(P["tareas"][t]["titulo"])))
        h.append('<div class="caja"><details open><summary>Enunciado que recibieron los modelos (versión zero-shot)</summary><pre>%s</pre></details></div>' % html.escape(P["tareas"][t]["enunciado"]))
        h.append('<div class="caja clave"><h3>Clave de evaluación</h3><ul>')
        if k["tipo"] == "diag":
            for ide, _, largo in k["items"]: h.append("<li><b>%s.</b> %s</li>" % (ide, html.escape(largo)))
            if k["extra"]: h.append("<li><i>%s</i></li>" % html.escape(k["extra"]))
        else:
            h.append("<li><b>Tabla de E/S.</b> Señales esperadas: " + "; ".join(html.escape(c.lower()) for c, _ in k["es"]) + ". Tipos de dato IEC 61131-3 válidos.</li>")
            h.append("<li><b>POU.</b> Tipo recomendado y si la justificación es coherente.</li>")
            for ide, _, largo in k["amb"]: h.append("<li><b>%s.</b> %s</li>" % (ide, html.escape(largo)))
        h.append("</ul></div>")
        h.append('<div class="caja idx"><b>Códigos, en orden:</b><br>' + "".join('<a href="#%s">%s</a>' % (x["codigo"], x["codigo"]) for x in resp) + "</div>")
        for i, x in enumerate(resp, 1):
            h.append('<div class="caja resp" id="%s"><div class="cod"><h3>%s</h3><span>%d de %d · %d caracteres</span></div><pre>%s</pre></div>' % (x["codigo"], x["codigo"], i, len(resp), len(x["texto"]), html.escape(x["texto"])))
    h.append("</main></body></html>")
    open(os.path.join(OUT, html_nombre), "w", encoding="utf-8").write("\n".join(h))
    return total

print(libro("Evaluador 1", False, "Ficha_Evaluador1.xlsx", "Respuestas_Evaluador1.html"))
print(libro("Evaluador 2", True, "Ficha_Evaluador2.xlsx", "Respuestas_Evaluador2.html"))
