# -*- coding: utf-8 -*-
"""Construye Fichas_Evaluacion_v2_720.xlsx: registro maestro de las 720 ejecuciones del experimento definitivo.

Lee (sin modificarlos) los resultados ya calculados:
  data/results/tabla_definitivo.csv            criterios 1 y 2 (CODESYS), 600 respuestas de código
  data/results/rubrica_definitivo.csv          rúbrica de 5 criterios; oficial = DeepSeek, seg_* = Claude (INC-35)
  data/results/evaluacion_manual_definitivo.csv  diagnóstico C1/C5/C6 y C2 (evaluador oficial DeepSeek)
  data/results/c4_definitivo.csv               verificación de C4
  src/tasks.json                               nivel, tipo y título de cada tarea
No calcula ninguna puntuación: solo reúne los datos y pone fórmulas de resumen en la hoja Métricas.
Uso (desde llm-plc-st-eval):  python analisis/construir_registro_maestro.py [carpeta_raiz_llm-plc-st-eval] [salida.xlsx]
"""
import csv, json, os, sys
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

RAIZ = sys.argv[1] if len(sys.argv) > 1 else os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
SALIDA = sys.argv[2] if len(sys.argv) > 2 else os.path.join(os.path.dirname(RAIZ), "02_Banco_y_Prompts", "Fichas_Evaluacion_v2_720.xlsx")
RES = os.path.join(RAIZ, "data", "results")
leer = lambda f: list(csv.DictReader(open(os.path.join(RES, f), encoding="utf-8-sig")))
ej = lambda e: str(e).replace("ejec_", "")

MODELOS = [("gemini-3.1-pro-preview", "Gemini 3.1 Pro"), ("claude-sonnet-4-6", "Claude Sonnet 4.6"),
           ("gpt-4o-2024-08-06", "GPT-4o"), ("qwen2.5-coder_14b", "Qwen2.5-Coder 14B")]
NOMBRE = dict(MODELOS)
ESTR = ["zero_shot", "few_shot"]
tareas = json.load(open(os.path.join(RAIZ, "src", "tasks.json"), encoding="utf-8"))["tareas"]

T = {(r["tarea"], r["modelo"], r["estrategia"], ej(r["ejec"])): r for r in leer("tabla_definitivo.csv")}
R = {(r["tarea"], r["modelo"], r["estrategia"], ej(r["ejec"])): r for r in leer("rubrica_definitivo.csv")}
M = {(r["tarea"], r["modelo"], r["estrategia"], ej(r["ejec"])): r for r in leer("evaluacion_manual_definitivo.csv")}
C4 = {("C4", r["modelo"], r["estrategia"], ej(r["ejec"])): r for r in leer("c4_definitivo.csv")}
assert len(T) == 600 and len(R) == 600 and len(M) == 96 and len(C4) == 24, (len(T), len(R), len(M), len(C4))

ARIAL = "Arial"
F = lambda **k: Font(name=ARIAL, size=k.pop("size", 10), **k)
GRIS = PatternFill("solid", fgColor="D9E1F2"); AZUL = PatternFill("solid", fgColor="DDEBF7"); AMAR = PatternFill("solid", fgColor="FFF2CC")
fino = Side(style="thin", color="BFBFBF"); BORDE = Border(left=fino, right=fino, top=fino, bottom=fino)

wb = Workbook()
# ------------------------------------------------------------------ Instructivo
ins = wb.active; ins.title = "Instructivo"
lineas = [
    ("REGISTRO MAESTRO DE EVALUACIÓN — EXPERIMENTO DEFINITIVO (720 EJECUCIONES)", F(bold=True, size=14)),
    ("Sustituye al registro de 624 ejecuciones de Fichas_Evaluacion.xlsx (piloto), que se conserva sin cambios.", F(italic=True)),
    ("", None),
    ("Qué contiene", F(bold=True, size=11)),
    ("Registro maestro: una fila por ejecución (30 tareas × 4 modelos × 2 estrategias × 3 ejecuciones = 720).", F()),
    ("Métricas: resúmenes con fórmulas sobre el registro (por modelo, por modelo y estrategia, por nivel y del segundo evaluador).", F()),
    ("", None),
    ("De dónde sale cada columna (archivos de llm-plc-st-eval/data/results/)", F(bold=True, size=11)),
    ("Compila, Pasa el harness, Categoría y Observaciones: tabla_definitivo.csv (CODESYS V3.5 SP22; criterios 2 y 1).", F()),
    ("Funcional, IEC, Coherencia, Claridad y Mantenibilidad: rubrica_definitivo.csv. Evaluador OFICIAL: DeepSeek deepseek-v4-pro (INC-35).", F()),
    ("  Reglas fijas: no compila → Funcional = 0 e IEC = 0; pasa todo el harness → Funcional = 3. El evaluador decide solo los niveles intermedios.", F()),
    ("Diagnóstico (C1, C5, C6) y C2: evaluacion_manual_definitivo.csv, mismo evaluador oficial.", F()),
    ("C4: c4_definitivo.csv (verificación automática de que el código quedó intacto).", F()),
    ("Segundo evaluador (Claude): solo para confiabilidad; no se usa en los resultados (columnas grises de la derecha).", F()),
    ("Código ciego: el código con que la respuesta se presentó a los evaluadores (asignación sellada).", F()),
    ("", None),
    ("Notas", F(bold=True, size=11)),
    ("Las celdas vacías significan 'no aplica' a ese tipo de tarea (por ejemplo, la rúbrica no se aplica a C1, C2, C4, C5 ni C6).", F()),
    ("Total (0–15) y Diagnóstico completo se calculan con fórmulas; el resto son datos copiados de los CSV, no se editan a mano.", F()),
    ("Si cambia algún CSV, se regenera el libro con analisis/construir_registro_maestro.py.", F()),
    ("Estadística inferencial (Wilson, McNemar, Wilcoxon, kappa): analisis/salidas/*.txt.", F()),
]
for i, (t, f) in enumerate(lineas, 1):
    c = ins.cell(i, 1, t)
    if f: c.font = f
ins.column_dimensions["A"].width = 140

# ------------------------------------------------------------------ Registro maestro
rm = wb.create_sheet("Registro maestro")
cols = ["ID", "Título", "Nivel", "Tipo", "Modelo", "Estrategia", "Ejec.", "Código ciego",
        "Compila (1/0)", "Pasa harness (1/0)", "Categoría criterio 1",
        "Funcional (0–3)", "IEC (0–3)", "Coherencia (0–3)", "Claridad (0–3)", "Mantenib. (0–3)", "Total (0–15)",
        "Detectados (0–4)", "Correcciones válidas (0–4)", "Hallazgos adicionales", "Errores inventados", "Diagnóstico completo (1/0)",
        "C2: E/S presentes (0–5)", "C2: ambigüedades clave (0–6)", "C2: POU", "C4: veredicto",
        "2.º eval. Funcional", "2.º eval. IEC", "2.º eval. Coherencia", "2.º eval. Claridad", "2.º eval. Mantenib.",
        "Observaciones / primer error"]
rm.cell(1, 1, "REGISTRO MAESTRO — 720 EJECUCIONES DEL EXPERIMENTO DEFINITIVO").font = F(bold=True, size=13)
rm.cell(2, 1, "Datos copiados de data/results/*.csv. Evaluador oficial de la rúbrica y del diagnóstico: DeepSeek deepseek-v4-pro. Columnas grises: segundo evaluador (Claude), solo confiabilidad.").font = F(italic=True)
H = 4
for j, n in enumerate(cols, 1):
    c = rm.cell(H, j, n); c.font = F(bold=True); c.fill = GRIS if j < 27 else PatternFill("solid", fgColor="E7E6E6")
    c.alignment = Alignment(wrap_text=True, vertical="center", horizontal="center"); c.border = BORDE
fila = H + 1
num = lambda v: int(v) if v not in ("", None) else None
for t in tareas:
    tid = t["id"]
    for mod, _ in MODELOS:
        for e in ESTR:
            for k in ("1", "2", "3"):
                key = (tid, mod, e, k)
                v = {c: None for c in cols}
                v.update({"ID": tid, "Título": t["titulo"], "Nivel": t["nivel"], "Tipo": t["tipo"], "Modelo": NOMBRE[mod],
                          "Estrategia": e, "Ejec.": int(k)})
                if key in T:
                    a, r = T[key], R[key]
                    v["Código ciego"] = r["codigo"]
                    v["Compila (1/0)"] = num(a["c2_compila"]); v["Pasa harness (1/0)"] = num(a["c1_pasa"])
                    v["Categoría criterio 1"] = a["c1_categoria"]
                    for col, k2 in (("Funcional (0–3)", "funcional"), ("IEC (0–3)", "iec"), ("Coherencia (0–3)", "coherencia"),
                                    ("Claridad (0–3)", "claridad"), ("Mantenib. (0–3)", "mantenibilidad")):
                        v[col] = num(r[k2])
                    for col, k2 in (("2.º eval. Funcional", "seg_funcional"), ("2.º eval. IEC", "seg_iec"), ("2.º eval. Coherencia", "seg_coherencia"),
                                    ("2.º eval. Claridad", "seg_claridad"), ("2.º eval. Mantenib.", "seg_mantenibilidad")):
                        v[col] = num(r[k2])
                    if a["c1_categoria"] == "NO_COMPILA":
                        v["Observaciones / primer error"] = (a["errores_texto"].split(" || ")[0])[:200]
                    elif a["c1_categoria"] == "FALLA_LOGICA":
                        v["Observaciones / primer error"] = "Pasos fallidos del harness: " + a["pasos_fallidos"]
                    elif a["c1_categoria"] == "INTERFAZ_NO_CONFORME":
                        v["Observaciones / primer error"] = "Compila, pero las E/S no son accesibles al harness"
                elif key in M:
                    m = M[key]; v["Código ciego"] = m["codigo"]
                    if tid == "C2":
                        v["C2: E/S presentes (0–5)"] = num(m["es_presentes"]); v["C2: ambigüedades clave (0–6)"] = num(m["ambiguedades_clave"])
                        v["C2: POU"] = m["pou"]
                    else:
                        v["Detectados (0–4)"] = num(m["detectados"]); v["Correcciones válidas (0–4)"] = num(m["correcciones_validas"])
                        v["Hallazgos adicionales"] = num(m["adicionales_validos"]); v["Errores inventados"] = num(m["errores_inventados"])
                    v["Observaciones / primer error"] = m["observaciones"]
                elif key in C4:
                    v["C4: veredicto"] = C4[key]["veredicto"]
                else:
                    raise SystemExit(f"Sin datos para {key}")
                for j, n in enumerate(cols, 1):
                    c = rm.cell(fila, j, v[n]); c.font = F(); c.border = BORDE
                    if j >= 27 and j <= 31: c.fill = PatternFill("solid", fgColor="F2F2F2")
                rm.cell(fila, 17, f'=IF(COUNT(L{fila}:P{fila})=5,SUM(L{fila}:P{fila}),"")').font = F(bold=True)
                rm.cell(fila, 22, f'=IF(R{fila}="","",IF(R{fila}=4,1,0))').font = F()
                fila += 1
ULT = fila - 1
assert ULT - H == 720, ULT - H
anchos = [6, 34, 13, 17, 18, 11, 6, 11, 9, 9, 22, 9, 8, 10, 9, 10, 9, 10, 11, 10, 10, 11, 10, 11, 16, 12, 9, 8, 10, 9, 10, 70]
for j, w in enumerate(anchos, 1): rm.column_dimensions[get_column_letter(j)].width = w
rm.row_dimensions[H].height = 45
rm.freeze_panes = rm.cell(H + 1, 3)
rm.auto_filter.ref = f"A{H}:{get_column_letter(len(cols))}{ULT}"

# ------------------------------------------------------------------ Métricas (fórmulas)
mt = wb.create_sheet("Métricas")
RG = lambda col: f"'Registro maestro'!${col}${H+1}:${col}${ULT}"
D, E_, Cn = RG("E"), RG("F"), RG("C")
mt.cell(1, 1, "MÉTRICAS — se calculan con fórmulas sobre el Registro maestro").font = F(bold=True, size=13)
mt.cell(2, 1, "Tasas como fracción (formato %). Medias de la rúbrica sobre las respuestas de código; del diagnóstico, sobre C1, C5 y C6.").font = F(italic=True)

def bloque(r0, titulo, filas, crit):
    """filas: lista de (etiqueta, [condiciones COUNTIFS como pares (rango, criterio)])"""
    mt.cell(r0, 1, titulo).font = F(bold=True, size=11)
    enc = ["Grupo", "Respuestas de código", "Compila", "Pasa harness", "Funcional", "IEC", "Coherencia", "Claridad", "Mantenib.", "Total (0–15)",
           "Detección diag.", "Diag. completo", "Inventados (media)"]
    for j, n in enumerate(enc, 1):
        c = mt.cell(r0 + 1, j, n); c.font = F(bold=True); c.fill = GRIS; c.border = BORDE; c.alignment = Alignment(wrap_text=True, horizontal="center")
    for i, (lab, conds) in enumerate(filas):
        rr = r0 + 2 + i
        cs = "".join(f",{rng},{crit_}" for rng, crit_ in conds)
        mt.cell(rr, 1, lab).font = F()
        f_ = {
            2: f'=COUNTIFS({RG("I")},"<>"{cs})',
            3: f'=IFERROR(AVERAGEIFS({RG("I")}{cs}),"")',
            4: f'=IFERROR(AVERAGEIFS({RG("J")}{cs}),"")',
            5: f'=IFERROR(AVERAGEIFS({RG("L")}{cs}),"")',
            6: f'=IFERROR(AVERAGEIFS({RG("M")}{cs}),"")',
            7: f'=IFERROR(AVERAGEIFS({RG("N")}{cs}),"")',
            8: f'=IFERROR(AVERAGEIFS({RG("O")}{cs}),"")',
            9: f'=IFERROR(AVERAGEIFS({RG("P")}{cs}),"")',
            10: f'=IFERROR(AVERAGEIFS({RG("Q")}{cs}),"")',
            11: f'=IFERROR(SUMIFS({RG("R")}{cs})/(4*COUNTIFS({RG("R")},"<>"{cs})),"")',
            12: f'=IFERROR(AVERAGEIFS({RG("V")}{cs}),"")',
            13: f'=IFERROR(AVERAGEIFS({RG("U")}{cs}),"")',
        }
        for j, fx in f_.items():
            c = mt.cell(rr, j, fx); c.font = F(); c.border = BORDE
            c.number_format = "0" if j == 2 else ("0.0%" if j in (3, 4, 11, 12) else "0.00")
    return r0 + 2 + len(filas) + 1

r = 4
r = bloque(r, "Por modelo", [(nm, [(D, f'"{nm}"')]) for _, nm in MODELOS] + [("Todos", [(RG("A"), '"<>"')])], None)
r = bloque(r, "Por modelo y estrategia", [(f"{nm} · {e}", [(D, f'"{nm}"'), (E_, f'"{e}"')]) for _, nm in MODELOS for e in ESTR], None)
r = bloque(r, "Por nivel", [(n, [(Cn, f'"{n}"')]) for n in ("Basico", "Intermedio", "Avanzado", "Complementaria")], None)

mt.cell(r, 1, "Segundo evaluador (Claude), solo confiabilidad: medias por modelo para comparar con el oficial").font = F(bold=True, size=11)
enc = ["Modelo", "Funcional", "IEC", "Coherencia", "Claridad", "Mantenib."]
for j, n in enumerate(enc, 1):
    c = mt.cell(r + 1, j, n); c.font = F(bold=True); c.fill = PatternFill("solid", fgColor="E7E6E6"); c.border = BORDE
for i, (_, nm) in enumerate(MODELOS):
    rr = r + 2 + i; mt.cell(rr, 1, nm).font = F()
    for j, col in enumerate("AA AB AC AD AE".split(), 2):
        c = mt.cell(rr, j, f'=IFERROR(AVERAGEIFS({RG(col)},{D},"{nm}"),"")'); c.font = F(); c.number_format = "0.00"; c.border = BORDE
mt.column_dimensions["A"].width = 34
for j in range(2, 14): mt.column_dimensions[get_column_letter(j)].width = 12
mt.freeze_panes = "B4"

os.makedirs(os.path.dirname(SALIDA), exist_ok=True)
wb.save(SALIDA)
print(f"Guardado {SALIDA}: {ULT - H} filas")
