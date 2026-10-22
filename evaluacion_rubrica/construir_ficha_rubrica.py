# -*- coding: utf-8 -*-
"""Construye las fichas Excel de la rubrica (5 criterios x 4 niveles) a partir de paquete_rubrica.json.
Una hoja por tarea; las puntuaciones que salen de CODESYS van precargadas y fijas.
Uso: python construir_ficha_rubrica.py <paquete_rubrica.json> <carpeta_salida>"""
import json, sys, os
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.comments import Comment
from openpyxl.utils import get_column_letter as L
P = json.load(open(sys.argv[1], encoding="utf-8")); OUT = sys.argv[2]; os.makedirs(OUT, exist_ok=True)
ORDEN = list(P["tareas"].keys())
F = lambda **k: Font(name="Arial", **({"size": 10} | k))
AMAR = PatternFill("solid", fgColor="FFF2CC"); GRIS = PatternFill("solid", fgColor="E7E6E6"); AZUL = PatternFill("solid", fgColor="1F4E78"); CALC = PatternFill("solid", fgColor="DDEBF7")
fino = Side(style="thin", color="BFBFBF"); BORDE = Border(left=fino, right=fino, top=fino, bottom=fino)
CEN = Alignment(horizontal="center", vertical="center", wrap_text=True); IZQ = Alignment(horizontal="left", vertical="center", wrap_text=True)
FILA_ENC, FILA_EJ, FILA_1 = 4, 5, 6
RUBRICA = [
 ("Corrección funcional", "No implementa la lógica pedida o falla en el caso nominal.", "Funciona en el caso nominal pero falla en casos límite evidentes (flancos, reset, E-stop).", "Correcto en el caso nominal y en la mayoría de los casos límite; fallas menores sin riesgo.", "Correcto en todos los casos probados, incluidos límites y secuencias de fallo.", "IEC 61508 / simulación CODESYS"),
 ("Cumplimiento IEC 61131-3", "No compila; sintaxis o declaraciones inválidas.", "Compila solo tras correcciones menores, o usa extensiones fuera del estándar.", "Compila sin cambios; declaraciones y tipos correctos.", "Compila sin cambios y usa construcciones idiomáticas del estándar (POU adecuado, tipos TIME correctos, FB estándar bien instanciados).", "IEC 61131-3 / compilador CODESYS"),
 ("Coherencia lógica", "Condiciones contradictorias, variables sin inicializar con efecto, o estados muertos que impiden operar.", "Lógica mayormente coherente con redundancias o condiciones muertas sin efecto grave.", "Sin contradicciones; inicialización correcta; transiciones completas.", "Además, robusta: interlocks completos, sin estados inalcanzables, comportamiento seguro ante entradas simultáneas.", "Análisis técnico + simulación"),
 ("Claridad estructural", "Sin estructura reconocible; nombres crípticos; sin comentarios.", "Estructura básica pero nombres pobres o comentarios ausentes/incorrectos.", "Nombres descriptivos, secciones ordenadas, comentarios útiles.", "Además, modularización adecuada (FB reutilizable, separación lógica/E/S) y comentarios que explican el porqué.", "Clean Code (Martin)"),
 ("Mantenibilidad", "Valores mágicos por doquier; imposible modificar sin reescribir.", "Modificable con esfuerzo; constantes parcialmente parametrizadas.", "Parámetros como variables/constantes declaradas; cambios localizados.", "Además, diseño extensible (parámetros de entrada del FB, sin duplicación de lógica).", "ISO/IEC 25010")]
COLS = [("Código", 11, None), ("Resultado en CODESYS (automático)", 46, "Viene de la compilación y de la simulación con el harness. No se edita."),
        ("Corrección funcional (0–3)", 13, "Precargada y fija si la respuesta pasa todos los pasos (3) o no compila (0). En amarillo: puntúa 0, 1 o 2 según la rúbrica, con el detalle de pasos fallidos del HTML."),
        ("Cumplimiento IEC (0–3)", 13, "Precargada y fija en 0 si no compila. En amarillo: 1 si usa extensiones fuera del estándar; 2 si compila sin cambios con declaraciones y tipos correctos; 3 si además usa construcciones idiomáticas."),
        ("Coherencia lógica (0–3)", 12, None), ("Claridad estructural (0–3)", 12, None), ("Mantenibilidad (0–3)", 13, None), ("Total (0–15)", 10, None),
        ("Justificación: corrección funcional / IEC", 34, "Una línea, solo para las celdas amarillas de esas dos columnas."),
        ("Justificación: coherencia", 30, None), ("Justificación: claridad", 30, None), ("Justificación: mantenibilidad", 30, None), ("Observaciones", 28, None), ("Estado", 11, None)]

def hoja(wb, t, resp):
    ws = wb.create_sheet(t); n = len(resp); ult = FILA_1 + n - 1
    ws["A1"] = "%s — %s (%s)" % (t, P["tareas"][t]["titulo"], P["tareas"][t]["nivel"]); ws["A1"].font = F(bold=True, size=13)
    ws["A2"] = "Lee cada respuesta por su código en respuestas/Respuestas_%s.html. Rellena solo las celdas amarillas. Las azules vienen de CODESYS o se calculan solas. La fila EJEMPLO no se analiza." % t; ws["A2"].font = F(italic=True)
    ws.row_dimensions[FILA_ENC].height = 48
    for j, (txt, ancho, nota) in enumerate(COLS, 1):
        c = ws.cell(row=FILA_ENC, column=j, value=txt); c.font = F(bold=True, color="FFFFFF"); c.fill = AZUL; c.alignment = CEN; c.border = BORDE
        ws.column_dimensions[L(j)].width = ancho
        if 5 <= j <= 7: nota = "%s — 0: %s 1: %s 2: %s 3: %s" % (RUBRICA[j - 3][0], RUBRICA[j - 3][1], RUBRICA[j - 3][2], RUBRICA[j - 3][3], RUBRICA[j - 3][4])
        if nota:
            cm = Comment(nota, "Ficha"); cm.width = 420; cm.height = 170; c.comment = cm
    d012 = DataValidation(type="list", formula1='"0,1,2"', allow_blank=True, showErrorMessage=True, errorTitle="Valor no válido", error="Aquí solo 0, 1 o 2: el 3 se reserva a las respuestas que pasan todos los pasos.")
    d123 = DataValidation(type="list", formula1='"1,2,3"', allow_blank=True, showErrorMessage=True, errorTitle="Valor no válido", error="Aquí solo 1, 2 o 3: el 0 se reserva a las respuestas que no compilan.")
    d0123 = DataValidation(type="list", formula1='"0,1,2,3"', allow_blank=True, showErrorMessage=True, errorTitle="Valor no válido", error="Escribe 0, 1, 2 o 3.")
    for d in (d012, d123, d0123): ws.add_data_validation(d)
    filas = [(FILA_EJ, {"codigo": "EJEMPLO", "resumen_codesys": "Compila. Falla 2 de 7 pasos del harness (pasos 3,5; numerados desde 0).", "func_fijo": None, "iec_fijo": None})] + [(FILA_1 + i, x) for i, x in enumerate(resp)]
    for r, x in filas:
        ej = (r == FILA_EJ); ws.row_dimensions[r].height = 30
        ws.cell(row=r, column=1, value=x["codigo"]); ws.cell(row=r, column=2, value=x["resumen_codesys"])
        if x["func_fijo"] is not None: ws.cell(row=r, column=3, value=x["func_fijo"])
        elif not ej: d012.add("C%d" % r)
        if x["iec_fijo"] is not None: ws.cell(row=r, column=4, value=x["iec_fijo"])
        elif not ej: d123.add("D%d" % r)
        if not ej: d0123.add("E%d:G%d" % (r, r))
        ws.cell(row=r, column=8, value='=IF(COUNT(C%d:G%d)=5,SUM(C%d:G%d),"")' % (r, r, r, r))
        ws.cell(row=r, column=14, value='=IF(COUNT(C%d:G%d)=5,"Completa","Falta")' % (r, r))
        if ej:
            for c, v in {3: 1, 4: 2, 5: 2, 6: 2, 7: 1, 9: "Arranca y para, pero rearranca solo tras la emergencia.", 10: "Sin contradicciones; transiciones completas.", 11: "Nombres claros, comentarios útiles.", 12: "El tiempo está escrito como literal en dos sitios."}.items(): ws.cell(row=r, column=c, value=v)
        for c in range(1, len(COLS) + 1):
            cel = ws.cell(row=r, column=c); cel.border = BORDE; cel.font = F(italic=ej, bold=(c == 1 and not ej))
            cel.alignment = IZQ if c in (2, 9, 10, 11, 12, 13) else Alignment(horizontal="center", vertical="center")
            if ej: cel.fill = GRIS
            elif c in (2, 8, 14): cel.fill = CALC
            elif c == 3: cel.fill = CALC if x["func_fijo"] is not None else AMAR
            elif c == 4: cel.fill = CALC if x["iec_fijo"] is not None else AMAR
            elif c in (5, 6, 7, 10, 11, 12, 13): cel.fill = AMAR
            elif c == 9: cel.fill = AMAR if (x["func_fijo"] is None or x["iec_fijo"] is None) else CALC
    ws.freeze_panes = "C%d" % FILA_1
    return ("N", FILA_1, ult, n)

def libro(nombre, solo, archivo):
    wb = Workbook(); ins = wb.active; ins.title = "Instrucciones"; info = {}
    for t in ORDEN:
        resp = [x for x in P["tareas"][t]["respuestas"] if x["evaluador2"] or not solo]; info[t] = hoja(wb, t, resp)
    total = sum(v[3] for v in info.values())
    txt = [
     ("Evaluación a ciegas con la rúbrica de 5 criterios — %s" % nombre, "t"),
     ("Tesis: LLM como asistentes de programación PLC en la academia (UTP). %d respuestas de código de 25 tareas (generación y conversión), cada una con un código que no revela el modelo ni la estrategia." % total, ""),
     ("", ""), ("Qué hay que hacer", "h"),
     ("1. Trabaja una tarea completa antes de pasar a la siguiente. Abre respuestas/Respuestas_<tarea>.html (enunciado, protocolo de prueba y respuestas por código) y la hoja de la misma tarea en este libro.", ""),
     ("2. Rellena ÚNICAMENTE las celdas amarillas con 0, 1, 2 o 3 según la hoja Rúbrica. Las celdas azules vienen de CODESYS o se calculan solas.", ""),
     ("3. Justifica cada puntuación en una línea, en la columna de justificación correspondiente. Sin justificación, la puntuación no es auditable.", ""),
     ("4. La fila gris EJEMPLO solo muestra el formato. No se analiza.", ""),
     ("5. Puedes repartir el trabajo en varias sesiones. La tabla de avance de abajo indica qué falta. Guarda con el mismo nombre.", ""),
     ("", ""), ("Reglas de decisión (fijadas antes de evaluar; no se cambian durante la evaluación)", "h"),
     ("Si la respuesta NO compila: Cumplimiento IEC = 0 y Corrección funcional = 0. Ya están precargadas. Los otros tres criterios se evalúan igual, leyendo el código.", ""),
     ("Si la respuesta pasa TODOS los pasos del harness: Corrección funcional = 3. Ya está precargada.", ""),
     ("Si compila pero falla pasos del harness: Corrección funcional es 0, 1 o 2 según la rúbrica. El HTML indica qué pasos fallaron y el protocolo de prueba dice qué comprueba cada paso.", ""),
     ("Si compila sola pero no se pudo simular (declara las entradas donde el harness no puede escribirlas): Corrección funcional es 0, 1 o 2, juzgada leyendo el código; no puede ser 3 porque no se demostró en simulación.", ""),
     ("Cumplimiento IEC cuando compila: 1 si usa extensiones fuera del estándar; 2 si las declaraciones y los tipos son correctos; 3 si además usa construcciones idiomáticas del estándar.", ""),
     ("Repartir el código entre declaración e implementación es un requisito del editor de CODESYS: no penaliza.", ""),
     ("Puntúa lo que el código es, no lo que el modelo probablemente quiso hacer.", ""),
     ("", ""), ("Evaluación a ciegas", "h"),
     ("No consultes _SELLADO_asignacion_rubrica.csv, la carpeta de datos ni los resultados por modelo mientras calificas. Los evaluadores trabajan por separado y no comentan casos concretos hasta que ambos terminen.", ""),
     ("", ""), ("Avance", "h")]
    for i, (s, e) in enumerate(txt, 1):
        c = ins.cell(row=i, column=1, value=s); c.alignment = Alignment(wrap_text=True, vertical="top")
        c.font = F(bold=True, size=14) if e == "t" else (F(bold=True, size=11, color="1F4E78") if e == "h" else F())
        ins.merge_cells(start_row=i, start_column=1, end_row=i, end_column=4)
        if len(s) > 105: ins.row_dimensions[i].height = 30
    f0 = len(txt) + 1
    for j, h in enumerate(["Tarea", "Respuestas", "Completas", "Pendientes"], 1):
        c = ins.cell(row=f0, column=j, value=h); c.font = F(bold=True, color="FFFFFF"); c.fill = AZUL; c.alignment = CEN; c.border = BORDE
    for i, t in enumerate(ORDEN, 1):
        col, a, b, n = info[t]; r = f0 + i
        ins.cell(row=r, column=1, value=t); ins.cell(row=r, column=2, value=n)
        ins.cell(row=r, column=3, value='=COUNTIF(%s!%s%d:%s%d,"Completa")' % (t, col, a, col, b)); ins.cell(row=r, column=4, value="=B%d-C%d" % (r, r))
        for j in range(1, 5):
            c = ins.cell(row=r, column=j); c.border = BORDE; c.font = F(); c.alignment = CEN
            if j >= 3: c.fill = CALC
    r = f0 + len(ORDEN) + 1
    ins.cell(row=r, column=1, value="Total")
    for j, col in ((2, "B"), (3, "C"), (4, "D")): ins.cell(row=r, column=j, value="=SUM(%s%d:%s%d)" % (col, f0 + 1, col, r - 1))
    for j in range(1, 5):
        c = ins.cell(row=r, column=j); c.border = BORDE; c.font = F(bold=True); c.alignment = CEN; c.fill = CALC
    for j, w in zip("ABCD", (46, 22, 22, 22)): ins.column_dimensions[j].width = w
    ins.cell(row=r + 2, column=1, value="Nombre del evaluador:").font = F(bold=True); ins.cell(row=r + 2, column=2).fill = AMAR
    ins.cell(row=r + 3, column=1, value="Fechas de calificación:").font = F(bold=True); ins.cell(row=r + 3, column=2).fill = AMAR
    ru = wb.create_sheet("Rúbrica", 1)
    ru["A1"] = "Rúbrica de evaluación — 5 criterios × 4 niveles"; ru["A1"].font = F(bold=True, size=13)
    ru["A2"] = "Fuente: Banco_de_Pruebas_v0.xlsx, hoja Rúbrica (2.5), sin cambios de redacción. Insuficiente = 0, Básico = 1, Competente = 2, Excelente = 3."; ru["A2"].font = F(italic=True)
    for j, h in enumerate(["Criterio", "Insuficiente (0)", "Básico (1)", "Competente (2)", "Excelente (3)", "Fuente del criterio"], 1):
        c = ru.cell(row=4, column=j, value=h); c.font = F(bold=True, color="FFFFFF"); c.fill = AZUL; c.alignment = CEN; c.border = BORDE
    for i, fila in enumerate(RUBRICA, 5):
        for j, v in enumerate(fila, 1):
            c = ru.cell(row=i, column=j, value=v); c.font = F(bold=(j == 1)); c.border = BORDE; c.alignment = Alignment(wrap_text=True, vertical="top")
        ru.row_dimensions[i].height = 78
    for j, w in zip("ABCDEF", (22, 34, 34, 34, 38, 22)): ru.column_dimensions[j].width = w
    wb.save(os.path.join(OUT, archivo)); return total
print(libro("Evaluador 1", False, "Ficha_Rubrica_Evaluador1.xlsx")); print(libro("Evaluador 2", True, "Ficha_Rubrica_Evaluador2.xlsx"))
