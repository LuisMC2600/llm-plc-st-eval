"""Compara una ficha humana (Evaluador1 o Evaluador2) con la ficha IA NO OFICIAL.
Empareja por código de respuesta. Solo usa filas humanas completas.

Uso:
  python comparar_humano_vs_ia.py diagnostico <ficha_humana.xlsx> <Ficha_IA_NO_OFICIAL_Diagnostico.xlsx> [salida.csv]
  python comparar_humano_vs_ia.py rubrica     <ficha_humana.xlsx> <Ficha_IA_NO_OFICIAL_Rubrica.xlsx>     [salida.csv]

Métricas: acuerdo exacto, kappa de Cohen (nominal) y, en la rúbrica, kappa ponderado
cuadrático, acuerdo ±1 y diferencia media (IA − humano). En la rúbrica, funcional e IEC se
reportan también solo sobre las celdas NO fijadas (las fijadas coinciden por construcción).
"""
import sys, csv
from collections import Counter
import openpyxl


def kappa(a, b, pesos=None, cats=None):
    cats = sorted(set(a) | set(b)) if cats is None else cats
    n = len(a)
    if n == 0:
        return float("nan")
    k = len(cats); idx = {c: i for i, c in enumerate(cats)}
    if pesos == "cuad":
        w = [[1 - ((i - j) / (k - 1)) ** 2 if k > 1 else 1 for j in range(k)] for i in range(k)]
    else:
        w = [[1.0 if i == j else 0.0 for j in range(k)] for i in range(k)]
    obs = [[0] * k for _ in range(k)]
    for x, y in zip(a, b):
        obs[idx[x]][idx[y]] += 1
    pa = [sum(r) / n for r in obs]; pb = [sum(obs[i][j] for i in range(k)) / n for j in range(k)]
    po = sum(w[i][j] * obs[i][j] / n for i in range(k) for j in range(k))
    pe = sum(w[i][j] * pa[i] * pb[j] for i in range(k) for j in range(k))
    return float("nan") if pe == 1 else (po - pe) / (1 - pe)


def leer(path, hojas, ncols):
    wb = openpyxl.load_workbook(path, data_only=False)
    out = {}
    for h in hojas:
        if h not in wb.sheetnames:
            continue
        ws = wb[h]; r = 6
        while ws.cell(r, 1).value:
            out[ws.cell(r, 1).value] = (h, [ws.cell(r, c).value for c in range(1, ncols + 1)], ws.cell(r, 2).value)
            r += 1
    return out


def norm(v):
    if v is None:
        return ""
    return str(v).strip().upper() if isinstance(v, str) else v


def resumen(nombre, a, b, filas, ordinal=False, cats=None):
    n = len(a)
    if n == 0:
        filas.append([nombre, 0] + [""] * 5); return
    ac = sum(x == y for x, y in zip(a, b)) / n
    fila = [nombre, n, f"{ac:.3f}", f"{kappa(a, b, cats=cats):.3f}"]
    if ordinal:
        fila += [f"{kappa(a, b, 'cuad', cats=cats):.3f}",
                 f"{sum(abs(x - y) <= 1 for x, y in zip(a, b)) / n:.3f}",
                 f"{sum(y - x for x, y in zip(a, b)) / n:+.3f}"]
    else:
        fila += ["", "", ""]
    filas.append(fila)


def main():
    modo, fh, fi = sys.argv[1:4]
    salida = sys.argv[4] if len(sys.argv) > 4 else f"comparacion_{modo}.csv"
    filas = []
    if modo == "diagnostico":
        H = leer(fh, ["C1", "C5", "C6", "C2"], 17); I = leer(fi, ["C1", "C5", "C6", "C2"], 17)
        comunes = [c for c in H if c in I]
        for t in ["C1", "C5", "C6"]:
            cs = [c for c in comunes if H[c][0] == t and all(H[c][1][j] not in (None, "") for j in (1, 3, 5, 7, 9, 10))]
            for e in range(4):
                resumen(f"{t} E{e+1} detectado", [norm(H[c][1][1 + 2*e]) for c in cs], [norm(I[c][1][1 + 2*e]) for c in cs], filas)
                cc = [c for c in cs if norm(H[c][1][1 + 2*e]) == "S" and norm(I[c][1][1 + 2*e]) == "S"]
                resumen(f"{t} E{e+1} corrección (ambos detectan)", [norm(H[c][1][2 + 2*e]) for c in cc], [norm(I[c][1][2 + 2*e]) for c in cc], filas)
            for j, nom in ((9, "adicionales"), (10, "inventados")):
                a = [int(H[c][1][j]) for c in cs]; b = [int(I[c][1][j]) for c in cs]
                resumen(f"{t} {nom} (n.º)", a, b, filas, ordinal=True, cats=sorted(set(a) | set(b)))
            a = [sum(norm(H[c][1][1 + 2*e]) == "S" for e in range(4)) for c in cs]
            b = [sum(norm(I[c][1][1 + 2*e]) == "S" for e in range(4)) for c in cs]
            resumen(f"{t} detectados (0–4)", a, b, filas, ordinal=True, cats=list(range(5)))
        cs = [c for c in comunes if H[c][0] == "C2" and all(H[c][1][j] not in (None, "") for j in range(1, 17))]
        nombres = ["E/S tanque", "E/S pozo", "E/S paro", "E/S bomba", "E/S luz", "tipos", "POU", "justificación",
                   "A1", "A2", "A3", "A4", "A5", "A6"]
        for j, nom in enumerate(nombres, start=1):
            resumen(f"C2 {nom}", [norm(H[c][1][j]) for c in cs], [norm(I[c][1][j]) for c in cs], filas)
        resumen("C2 código incluido", [norm(H[c][1][16]) for c in cs], [norm(I[c][1][16]) for c in cs], filas)
        a = [int(H[c][1][15]) for c in cs]; b = [int(I[c][1][15]) for c in cs]
        resumen("C2 ambigüedades adicionales (n.º)", a, b, filas, ordinal=True, cats=sorted(set(a) | set(b)))
    elif modo == "rubrica":
        tareas = ["B1","B2","B3","B4","B5","B6","B7","B8","I1","I2","I3","I4","I5","I6","I7","I8",
                  "A1","A2","A3","A4","A5","A6","C3","C7","C8"]
        H = leer(fh, tareas, 7); I = leer(fi, tareas, 7)
        cs = [c for c in H if c in I and all(isinstance(H[c][1][j], (int, float)) for j in range(2, 7))]
        crit = ["funcional", "iec", "coherencia", "claridad", "mantenibilidad"]
        def fijado(c, k):
            res = str(H[c][2] or "")
            if res.startswith("No compila"): return True
            return k == "funcional" and "Pasa" in res and "Falla" not in res
        for nivel, grupo in [("Global", None), ("Básico", "B"), ("Intermedio", "I"), ("Avanzado", "A"), ("Comparativa", "C")]:
            sub = [c for c in cs if grupo is None or H[c][0].startswith(grupo)]
            for j, k in enumerate(crit, start=2):
                resumen(f"{nivel} {k}", [int(H[c][1][j]) for c in sub], [int(I[c][1][j]) for c in sub], filas, True, list(range(4)))
                if k in ("funcional", "iec"):
                    s2 = [c for c in sub if not fijado(c, k)]
                    resumen(f"{nivel} {k} (solo no fijadas)", [int(H[c][1][j]) for c in s2], [int(I[c][1][j]) for c in s2], filas, True, list(range(4)))
            a = [sum(int(H[c][1][j]) for j in range(2, 7)) for c in sub]
            b = [sum(int(I[c][1][j]) for j in range(2, 7)) for c in sub]
            resumen(f"{nivel} total (0–15)", a, b, filas, True, list(range(16)))
        # desacuerdos grandes para revisión
        with open(salida.replace(".csv", "_desacuerdos.csv"), "w", newline="", encoding="utf-8") as f:
            w = csv.writer(f); w.writerow(["codigo", "criterio", "humano", "ia"])
            for c in cs:
                for j, k in enumerate(crit, start=2):
                    if abs(int(H[c][1][j]) - int(I[c][1][j])) >= 2:
                        w.writerow([c, k, H[c][1][j], I[c][1][j]])
    else:
        sys.exit("modo debe ser diagnostico o rubrica")
    with open(salida, "w", newline="", encoding="utf-8") as f:
        w = csv.writer(f)
        w.writerow(["item", "n", "acuerdo", "kappa", "kappa_pond_cuad", "acuerdo_pm1", "dif_media_IA_menos_humano"])
        w.writerows(filas)
    for fila in filas:
        print(" | ".join(str(x) for x in fila))
    print(f"\nGuardado en {salida}")


if __name__ == "__main__":
    main()
